"""Production application form: render, CV generation and webhook delivery.

Uses fictional applicants and mocks delivery; never writes to the live Sheet.
"""
import base64
import io
import json
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from unittest.mock import patch

from django.core.cache import cache
from django.test import Client, SimpleTestCase, override_settings
from docx import Document
from PIL import Image

from aimtravel_site.web import application_form_view as form


@override_settings(ALLOWED_HOSTS=["testserver"])
class PublicApplicationFormTest(SimpleTestCase):
    def setUp(self):
        cache.clear()

    def payload(self):
        values = [""] * 31
        values[0] = "TEST APPLICATION"
        values[1] = "2005-06-14"
        values[5] = "+359888000000"
        values[6] = "test.applicant@example.com"
        values[13] = "Test University"
        values[21] = "2027-05-15"
        photo = io.BytesIO()
        Image.new("RGB", (8, 8), "white").save(photo, format="PNG")
        return {"values": values, "consent": True,
                "submissionId": "test-submission-1234567890",
                "photo": base64.b64encode(photo.getvalue()).decode("ascii"),
                "gender": "Female", "nationality": "Bulgaria", "experiences": []}

    def submit(self, body):
        return self.client.post("/application-form/submit/", body,
                                content_type="application/json")

    def test_page_and_csrf(self):
        client = Client(enforce_csrf_checks=True)
        response = client.get("/application-form/")
        self.assertContains(response, "Application Form")
        self.assertContains(response, "/application-form/submit/")
        self.assertContains(response, 'AIM_logo-ROUNDED.png')
        self.assertContains(response, 'aria-describedby="photo-hint photo-error" required')
        self.assertIn("csrftoken", response.cookies)
        self.assertEqual(client.post("/application-form/submit/", self.payload(),
                                    content_type="application/json").status_code, 403)

    def test_fresh_session_token_keeps_csrf_protection_enabled(self):
        client = Client(enforce_csrf_checks=True)
        response = client.get('/application-form/session/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('no-store', response['Cache-Control'])
        token = response.json()['csrfToken']
        self.assertEqual(len(token), 64)
        self.assertEqual(client.post('/application-form/submit/', self.payload(), content_type='application/json').status_code, 403)
        payload = self.payload()
        payload['consent'] = False
        accepted = client.post('/application-form/submit/', payload, content_type='application/json', HTTP_X_CSRFTOKEN=token)
        self.assertEqual(accepted.status_code, 400)
        self.assertIn('agree', accepted.json()['error'])

    def test_validation_identifies_the_field_and_form_is_not_cached(self):
        self.assertIn('no-store', self.client.get('/application-form/')['Cache-Control'])
        payload = self.payload()
        payload['values'][3] = 'София'
        result = self.submit(payload)
        self.assertEqual(result.status_code, 400)
        self.assertIn('address', result.json()['fields'])

    def test_cv_is_generated_from_the_shipped_template(self):
        payload = self.payload()
        document = Document(io.BytesIO(form._build_resume(
            payload["values"], [], payload["photo"], "Female", "Bulgaria")))
        self.assertGreater(len(document.inline_shapes), 0)
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n" + "\n".join(c.text for c in form._iter_cells(document.tables))
        self.assertIn("TEST APPLICATION", text)
        self.assertIn("test.applicant@example.com", text)
        self.assertIn("14/06/2005", text)
        self.assertIn("15/05/2027", text)
        self.assertNotIn("[FULL NAME IN LATIN LETTERS]", text)

    @patch.object(form, "APPLICATION_FORM_WEBHOOK_URL", "https://script.google.com/macros/s/test/exec")
    @patch.object(form, "APPLICATION_FORM_WEBHOOK_SECRET", "test-only-secret")
    @patch.object(form.urllib.request, "urlopen")
    def test_submission_preserves_sheet_payload_and_cv_and_does_not_duplicate(self, outbound):
        outbound.return_value.__enter__.return_value.read.return_value = b'{"ok":true}'
        payload = self.payload()
        self.assertEqual(self.submit(payload).json(), {"ok": True})
        request = outbound.call_args.args[0]
        sent = json.loads(request.data)
        self.assertEqual(sent["values"], payload["values"])
        self.assertEqual(sent["submissionId"], payload["submissionId"])
        self.assertEqual(sent["secret"], "test-only-secret")
        self.assertEqual(sent["resumeFileName"], "TEST_APPLICATION_AIM_CV.docx")
        self.assertTrue(base64.b64decode(sent["resumeBase64"]).startswith(b"PK"))
        self.assertEqual(self.submit(payload).json(), {"ok": True})
        self.assertEqual(outbound.call_count, 1)

    @patch.object(form.urllib.request, "urlopen")
    def test_invalid_input_never_reaches_google(self, outbound):
        payload = self.payload()
        payload["consent"] = False
        self.assertEqual(self.submit(payload).status_code, 400)
        payload["consent"] = True
        payload["values"][0] = "Тест"
        self.assertEqual(self.submit(payload).status_code, 400)
        outbound.assert_not_called()

    @patch.object(form, "APPLICATION_FORM_WEBHOOK_URL", "")
    def test_unconfigured_form_reports_unavailable(self):
        self.assertEqual(self.submit(self.payload()).status_code, 503)

    @patch.object(form.urllib.request, "urlopen")
    def test_photo_is_required_and_must_be_a_real_image(self, outbound):
        for photo in ("", "not-base64", base64.b64encode(b"not an image").decode()):
            with self.subTest(photo=photo):
                payload = self.payload()
                payload["photo"] = photo
                self.assertEqual(self.submit(payload).status_code, 400)
        outbound.assert_not_called()

    @patch.object(form, "APPLICATION_FORM_WEBHOOK_URL", "https://script.google.com/macros/s/test/exec")
    @patch.object(form, "APPLICATION_FORM_WEBHOOK_SECRET", "test-only-secret")
    @patch.object(form.urllib.request, "urlopen")
    def test_two_applicants_on_same_ip_are_not_rate_limited(self, outbound):
        outbound.return_value.__enter__.return_value.read.return_value = b'{"ok":true}'
        first = self.payload()
        second = self.payload()
        second["values"][6] = "second.applicant@example.com"
        second["submissionId"] = "second-submission-1234567890"
        self.assertEqual(self.submit(first).status_code, 200)
        self.assertEqual(self.submit(second).status_code, 200)
        self.assertEqual(outbound.call_count, 2)

    @patch.object(form, "APPLICATION_FORM_WEBHOOK_URL", "https://script.google.com/macros/s/test/exec")
    @patch.object(form, "APPLICATION_FORM_WEBHOOK_SECRET", "test-only-secret")
    @patch.object(form.urllib.request, "urlopen")
    def test_two_simultaneous_applicants_keep_their_own_payloads(self, outbound):
        barrier = Barrier(2)
        sent = []

        def delivery(request, timeout):
            sent.append(json.loads(request.data))
            barrier.wait(timeout=10)
            return io.BytesIO(b'{"ok":true}')

        outbound.side_effect = delivery
        first = self.payload()
        second = self.payload()
        second["values"][0] = "SECOND TEST APPLICATION"
        second["values"][6] = "second.applicant@example.com"
        second["submissionId"] = "second-submission-1234567890"

        def submit(payload):
            return Client().post("/application-form/submit/", payload,
                                 content_type="application/json")

        with ThreadPoolExecutor(max_workers=2) as workers:
            responses = list(workers.map(submit, [first, second]))
        self.assertEqual([response.status_code for response in responses], [200, 200])
        self.assertEqual({item["submissionId"] for item in sent},
                         {first["submissionId"], second["submissionId"]})
        self.assertEqual({item["values"][6] for item in sent},
                         {first["values"][6], second["values"][6]})
        self.assertEqual(len({item["resumeBase64"] for item in sent}), 2)

    @patch.object(form, "APPLICATION_FORM_WEBHOOK_URL", "https://script.google.com/macros/s/test/exec")
    @patch.object(form, "APPLICATION_FORM_WEBHOOK_SECRET", "test-only-secret")
    @patch.object(form.urllib.request, "urlopen")
    def test_busy_webhook_is_retryable_with_same_submission_id(self, outbound):
        outbound.return_value.__enter__.return_value.read.side_effect = [
            b'{"ok":false,"error":"busy"}', b'{"ok":true}']
        payload = self.payload()
        self.assertEqual(self.submit(payload).status_code, 503)
        self.assertEqual(self.submit(payload).json(), {"ok": True})
        self.assertEqual(outbound.call_count, 2)
