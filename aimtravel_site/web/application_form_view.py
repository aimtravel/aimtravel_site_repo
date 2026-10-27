"""Application form and CV delivery through the existing Google Apps Script webhook."""

import hashlib
import base64
import binascii
import io
import json
import re
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

from django.core.cache import cache
from django.http import JsonResponse
from django.shortcuts import render
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_GET, require_POST

try:
    from application_form_secrets import (
        APPLICATION_FORM_WEBHOOK_SECRET,
        APPLICATION_FORM_WEBHOOK_URL,
    )
except ImportError:
    APPLICATION_FORM_WEBHOOK_SECRET = ""
    APPLICATION_FORM_WEBHOOK_URL = ""


EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
FIELD_COUNT = 31
FIELD_LIMIT = 5000
MAX_REQUEST_BYTES = 6 * 1024 * 1024
MAX_PHOTO_BYTES = 3 * 1024 * 1024
RESUME_TEMPLATE = Path(__file__).resolve().parents[1] / "assets" / "AIM_Travel_USA_Modern_Resume_Template.docx"


def _format_date(value):
    value = str(value or "").strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        year, month, day = value.split("-")
        return f"{day}/{month}/{year}"
    return value


def _format_work_period(start_date, end_date, legacy_period=""):
    def format_month(value):
        value = str(value or "").strip()
        if re.fullmatch(r"\d{4}-\d{2}", value):
            year, month = value.split("-")
            return f"{month}/{year}"
        return _format_date(value)

    start = format_month(start_date)
    end = format_month(end_date)
    if start and end:
        return f"{start} – {end}"
    return start or end or str(legacy_period or "").strip()


def _replace_paragraph(paragraph, value):
    value = str(value or "")
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _iter_cells(tables, seen=None):
    if seen is None:
        seen = set()
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                # Merged table cells are exposed once per covered column by python-docx.
                # Deduplicate them so each CV placeholder consumes exactly one experience.
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                yield cell
                yield from _iter_cells(cell.tables, seen)


def _build_resume(values, experiences, photo_data, gender, nationality):
    """Fill the approved one-page CV template in memory; exclude application-only data."""
    if not RESUME_TEMPLATE.is_file():
        raise FileNotFoundError("CV template is not installed")

    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    document = Document(str(RESUME_TEMPLATE))
    public_fields = {
        "GENDER": gender,
        "DATE OF BIRTH": _format_date(values[1]),
        "NATIONALITY": nationality,
        "AVAILABLE START DATE": _format_date(values[21]),
        "PHONE NUMBER": values[5],
        "E-MAIL": values[6],
        "MICROSOFT TEAMS": values[8],
        "INSTAGRAM": values[7],
        "UNIVERSITY": values[13],
        "MAJOR / FIELD OF STUDY": values[14],
        "YEAR OF STUDIES": values[15],
        "ENGLISH SKILLS (1–10)": values[24],
        "OTHER LANGUAGES": values[25],
        "PERSONAL QUALITIES": ", ".join(filter(None, (values[26], values[27]))),
        "HOBBIES": ", ".join(filter(None, (values[28], values[29]))),
    }
    for paragraph in document.paragraphs:
        if "[FULL NAME IN LATIN LETTERS]" in paragraph.text:
            _replace_paragraph(paragraph, values[0])

    experience_fields = {
        "DATES": [_format_work_period(item.get("start_date", ""), item.get("end_date", ""), item.get("period", "")) for item in experiences],
        "POSITION HELD": [item.get("position", "") for item in experiences],
        "EMPLOYER / WORKPLACE": [item.get("employer", "") for item in experiences],
        "BUSINESS SECTOR": [", ".join(filter(None, (item.get("city", ""), item.get("country", "")))) for item in experiences],
        "MAIN ACTIVITIES AND RESPONSIBILITIES": [item.get("description", "") for item in experiences],
    }
    seen = {label: 0 for label in experience_fields}
    for cell in _iter_cells(document.tables):
        if not cell.paragraphs:
            continue
        label = cell.paragraphs[0].text.strip().upper()
        if label in public_fields and len(cell.paragraphs) > 1:
            _replace_paragraph(cell.paragraphs[1], public_fields[label])
        elif label in experience_fields and len(cell.paragraphs) > 1:
            index = seen[label]
            items = experience_fields[label]
            _replace_paragraph(cell.paragraphs[1], items[index] if index < len(items) else "")
            seen[label] += 1
            if label == "BUSINESS SECTOR":
                _replace_paragraph(cell.paragraphs[0], "CITY / COUNTRY")

    if not photo_data:
        raise ValueError("A photo is required")
    if photo_data:
        try:
            photo_bytes = base64.b64decode(photo_data, validate=True)
        except (ValueError, binascii.Error):
            raise ValueError("Invalid photo data")
        if not photo_bytes or len(photo_bytes) > MAX_PHOTO_BYTES:
            raise ValueError("Photo is too large")
        from PIL import Image, UnidentifiedImageError
        try:
            with Image.open(io.BytesIO(photo_bytes)) as image:
                if image.format not in {"JPEG", "PNG", "WEBP"}:
                    raise ValueError("Unsupported photo format")
                image.verify()
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
            raise ValueError("Invalid photo") from exc
        photo_cell = document.tables[0].rows[0].cells[2]
        photo_cell.text = ""
        photo_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = photo_cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(io.BytesIO(photo_bytes), height=Inches(0.92))

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _experience_items(raw_items):
    if raw_items is None:
        return []
    if not isinstance(raw_items, list) or len(raw_items) > 2:
        raise ValueError("Invalid experience list")
    result = []
    for item in raw_items:
        if not isinstance(item, dict):
            raise ValueError("Invalid experience entry")
        entry = {key: str(item.get(key, "") or "").strip() for key in ("position", "employer", "country", "city", "start_date", "end_date", "period", "description")}
        if any(len(value) > FIELD_LIMIT for value in entry.values()):
            raise ValueError("Experience entry is too long")
        if any(entry.values()):
            result.append(entry)
    return result


@require_GET
def application_form_page(request):
    return render(request, "application-form-mockup.html")


@require_POST
@csrf_protect
def application_form_submit(request):
    if len(request.body) > MAX_REQUEST_BYTES:
        return JsonResponse({"ok": False, "error": "The uploaded file is too large."}, status=413)
    try:
        body = json.loads(request.body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return JsonResponse({"ok": False, "error": "Invalid request."}, status=400)

    if not isinstance(body, dict):
        return JsonResponse({"ok": False, "error": "Invalid request."}, status=400)

    # Bots that fill the visually-hidden field are silently discarded.
    if str(body.get("website", "")).strip():
        return JsonResponse({"ok": True})

    if body.get("consent") is not True:
        return JsonResponse({"ok": False, "error": "Please confirm that you agree to the processing of your information."}, status=400)

    submission_id = str(body.get("submissionId", ""))
    if not re.fullmatch(r"[A-Za-z0-9_-]{16,80}", submission_id):
        return JsonResponse({"ok": False, "error": "This request is invalid. Refresh the page and try again."}, status=400)

    raw_values = body.get("values")
    if not isinstance(raw_values, list) or len(raw_values) != FIELD_COUNT:
        return JsonResponse({"ok": False, "error": "Please check your answers and try again."}, status=400)

    values = []
    for item in raw_values:
        if item is None:
            value = ""
        elif isinstance(item, (str, int, float, bool)):
            value = str(item).strip()
        else:
            return JsonResponse({"ok": False, "error": "Please check your answers and try again."}, status=400)
        if len(value) > FIELD_LIMIT:
            return JsonResponse({"ok": False, "error": "One of the fields is too long."}, status=400)
        if not value.isascii():
            return JsonResponse({"ok": False, "error": "Please use English and Latin characters only."}, status=400)
        values.append(value)

    full_name, email = values[0], values[6].lower()
    if not full_name or not email or not EMAIL_RE.match(email):
        return JsonResponse({"ok": False, "error": "Enter your full name and a valid email address."}, status=400)
    values[0] = full_name
    values[6] = email

    gender = str(body.get("gender", "") or "").strip()
    nationality = str(body.get("nationality", "") or "").strip()
    if (gender and gender not in {"Female", "Male", "Other", "Prefer not to say"}) or (nationality and not re.fullmatch(r"[A-Za-z][A-Za-z .'-]{0,59}", nationality)):
        return JsonResponse({"ok": False, "error": "Please select a valid gender and nationality."}, status=400)

    try:
        experiences = _experience_items(body.get("experiences", []))
        if any(not value.isascii() for entry in experiences for value in entry.values()):
            raise ValueError("Latin characters only")
        photo_data = str(body.get("photo", "") or "")
        if not photo_data:
            return JsonResponse({"ok": False, "error": "Please add a photo for your CV."}, status=400)
        if len(photo_data) > (MAX_PHOTO_BYTES * 4 // 3 + 16):
            raise ValueError("Photo is too large")
        resume_bytes = _build_resume(values, experiences, photo_data, gender, nationality)
    except (ValueError, FileNotFoundError):
        return JsonResponse({"ok": False, "error": "We could not prepare your CV. Check the photo and try again."}, status=400)
    except Exception:
        return JsonResponse({"ok": False, "error": "The CV service is temporarily unavailable. Please try again later."}, status=503)

    if not APPLICATION_FORM_WEBHOOK_URL or not APPLICATION_FORM_WEBHOOK_SECRET:
        return JsonResponse({"ok": False, "error": "The form is temporarily unavailable. Please try again later."}, status=503)

    # Different applicants can share an office/Wi-Fi IP. Only throttle repeat
    # delivery for this email; retain the same submission ID on failed retries.
    email_digest = hashlib.sha256(email.encode("utf-8")).hexdigest()
    email_cache_key = "application_form_email_" + email_digest
    submission_cache_key = "application_form_submission_" + submission_id
    submission_status = cache.get(submission_cache_key)
    if submission_status == "done":
        return JsonResponse({"ok": True})
    if cache.get(email_cache_key):
        return JsonResponse({"ok": False, "error": "An application with this email was recently submitted. Please wait a moment."}, status=429)
    if submission_status == "in_progress":
        return JsonResponse({"ok": False, "error": "This application is already being submitted. Please wait a moment."}, status=409)
    if not cache.add(submission_cache_key, "in_progress", timeout=90):
        return JsonResponse({"ok": False, "error": "This application is already being submitted. Please wait a moment."}, status=409)

    endpoint = urllib.parse.urlparse(APPLICATION_FORM_WEBHOOK_URL)
    if endpoint.scheme != "https" or endpoint.hostname != "script.google.com":
        cache.delete(submission_cache_key)
        return JsonResponse({"ok": False, "error": "The form is temporarily unavailable. Please try again later."}, status=503)

    payload = json.dumps(
        {
            "secret": APPLICATION_FORM_WEBHOOK_SECRET,
            "website": "",
            "submissionId": submission_id,
            "values": values,
            "gender": gender,
            "nationality": nationality,
            "resumeFileName": re.sub(r"[^A-Za-z0-9._-]+", "_", full_name).strip("._-")[:80] + "_AIM_CV.docx",
            "resumeBase64": base64.b64encode(resume_bytes).decode("ascii"),
        },
        ensure_ascii=False,
    ).encode("utf-8")
    outbound = urllib.request.Request(
        APPLICATION_FORM_WEBHOOK_URL,
        data=payload,
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    try:
        with urllib.request.urlopen(outbound, timeout=60) as response:
            result = json.loads(response.read(4096).decode("utf-8"))
    except (urllib.error.URLError, TimeoutError, UnicodeDecodeError, json.JSONDecodeError):
        cache.delete(submission_cache_key)
        return JsonResponse({"ok": False, "error": "We could not submit your application. Please try again shortly."}, status=502)

    if not isinstance(result, dict) or result.get("ok") is not True:
        cache.delete(submission_cache_key)
        if isinstance(result, dict) and result.get("error") == "busy":
            return JsonResponse({"ok": False, "error": "Other applications are being saved. Please wait a moment and submit again. Your answers are still here."}, status=503)
        if isinstance(result, dict) and result.get("error") == "rate_limited":
            return JsonResponse({"ok": False, "error": "rate_limited"}, status=429)
        return JsonResponse({"ok": False, "error": "We could not submit your application. Please try again shortly."}, status=502)

    cache.set(email_cache_key, "1", timeout=60)
    cache.set(submission_cache_key, "done", timeout=86400)
    return JsonResponse({"ok": True})
